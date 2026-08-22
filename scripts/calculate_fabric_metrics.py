import json
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "raw_data" / "phase1_3" / "Fabric_sensor_tests_Origin_simple.xlsx"
OUTPUT = ROOT / "generated" / "fabric"
OUTPUT.mkdir(parents=True, exist_ok=True)
OUT_JSON = OUTPUT / "fabric_metrics.json"
OUT_NPZ = OUTPUT / "fabric_records.npz"

BLOCKS = {
    "Grid80_standard": (1, 4),
    "Grid80_concave": (6, 9),
    "Rect80_standard": (11, 14),
    "Rect80_concave": (16, 19),
    "Rect100_convex": (21, 24),
    "Rect100_concave": (26, 29),
}


def f(value):
    if value is None:
        return np.nan
    return float(value)


wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb["Origin_simple"]

records = {}
results = []
npz_payload = {}

for label, (start_col, end_col) in BLOCKS.items():
    rows = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row,
                            min_col=start_col, max_col=end_col,
                            values_only=True):
        cycle, time_s, resistance, normalized = map(f, row)
        if np.isfinite(cycle) and np.isfinite(normalized):
            rows.append((cycle, time_s, resistance, normalized))
    arr = np.asarray(rows, dtype=float)
    cycle_number = np.floor(arr[:, 0] + 1e-8).astype(int)
    valid_cycle_ids = sorted(set(cycle_number.tolist()))
    cycle_metrics = []
    for cycle_id in valid_cycle_ids:
        values = arr[cycle_number == cycle_id, 3]
        cycle_metrics.append({
            "cycle": cycle_id,
            "n_observations": int(values.size),
            "center_pct": float(np.median(values)),
            "a80_pct_points": float(np.percentile(values, 90) - np.percentile(values, 10)),
        })

    early = [x for x in cycle_metrics if 1 <= x["cycle"] <= 50]
    late = [x for x in cycle_metrics if 451 <= x["cycle"] <= 500]
    late_amplitudes = np.array([x["a80_pct_points"] for x in late], dtype=float)
    evolution = np.median([x["center_pct"] for x in late]) - np.median(
        [x["center_pct"] for x in early]
    )
    late_cv = 100 * np.std(late_amplitudes, ddof=1) / np.mean(late_amplitudes)

    counts = np.array([x["n_observations"] for x in cycle_metrics])
    result = {
        "record": label,
        "n_rows": int(arr.shape[0]),
        "first_cycle_value": float(arr[0, 0]),
        "last_cycle_value": float(arr[-1, 0]),
        "cycle_count": len(valid_cycle_ids),
        "first_cycle": int(min(valid_cycle_ids)),
        "last_cycle": int(max(valid_cycle_ids)),
        "min_points_per_cycle": int(counts.min()),
        "median_points_per_cycle": float(np.median(counts)),
        "max_points_per_cycle": int(counts.max()),
        "early_median_a80_pct_points": float(np.median([x["a80_pct_points"] for x in early])),
        "late_median_a80_pct_points": float(np.median(late_amplitudes)),
        "all_cycle_median_a80_pct_points": float(np.median([x["a80_pct_points"] for x in cycle_metrics])),
        "resistance_evolution_pct_points": float(evolution),
        "late_amplitude_cv_pct": float(late_cv),
        "early_center_pct": float(np.median([x["center_pct"] for x in early])),
        "late_center_pct": float(np.median([x["center_pct"] for x in late])),
        "raw_min_pct": float(np.min(arr[:, 3])),
        "raw_max_pct": float(np.max(arr[:, 3])),
        "cycle_metrics": cycle_metrics,
    }
    results.append(result)
    records[label] = arr
    npz_payload[f"{label}_data"] = arr
    npz_payload[f"{label}_cycle"] = np.array([x["cycle"] for x in cycle_metrics])
    npz_payload[f"{label}_center"] = np.array([x["center_pct"] for x in cycle_metrics])
    npz_payload[f"{label}_a80"] = np.array([x["a80_pct_points"] for x in cycle_metrics])

OUT_JSON.write_text(json.dumps(results, indent=2))
np.savez_compressed(OUT_NPZ, **npz_payload)

for result in results:
    print(
        result["record"],
        f"rows={result['n_rows']}",
        f"cycles={result['first_cycle']}-{result['last_cycle']} ({result['cycle_count']})",
        f"points/cycle={result['min_points_per_cycle']}/{result['median_points_per_cycle']:.0f}/{result['max_points_per_cycle']}",
        f"A80_all={result['all_cycle_median_a80_pct_points']:.3f}",
        f"A80_late={result['late_median_a80_pct_points']:.3f}",
        f"evolution={result['resistance_evolution_pct_points']:.3f}",
        f"late_CV={result['late_amplitude_cv_pct']:.3f}",
        f"range=[{result['raw_min_pct']:.2f},{result['raw_max_pct']:.2f}]",
    )
